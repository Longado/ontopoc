"""Turn an uploaded business table (CSV or Excel) into the source bundle the ontology builder reads."""
from __future__ import annotations

import csv
from datetime import date, datetime
import hashlib
import io
from pathlib import PurePath
import zipfile

MAX_BYTES = 10 * 1024 * 1024  # ponytail: whole file in memory; stream rows if customers bring bigger exports
DEFAULT_PURPOSE = '描述这家公司的业务：有哪些对象（如客户、产品、订单）、它们之间怎样关联，让管理者的业务问题能在数据上回答。'
TABLE_SUFFIXES = ('.csv', '.xlsx')


class SourceFileError(ValueError):
    """The uploaded file cannot be read as business data; the message is shown to the user."""


def _cell(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(timespec='seconds')
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _headers(row) -> list[str]:
    names, seen = [], {}
    for i, cell in enumerate(row, start=1):
        name = _cell(cell) or f'列{i}'
        seen[name] = seen.get(name, 0) + 1
        names.append(name if seen[name] == 1 else f'{name}_{seen[name]}')
    return names


def _width(row) -> int:
    return sum(_cell(c) is not None for c in row)


def _records(rows: list[list]) -> tuple[list[dict], list[str]]:
    """Rows to records. Title lines above the real header (fewer filled cells than half the widest row) are skipped
    and returned, so exports that start with a report title still read their real header."""
    rows = [r for r in rows if _width(r) > 0]
    if not rows:
        return [], []
    widest = max(_width(r) for r in rows[:10])
    needed = max(2, (widest + 1) // 2) if widest >= 2 else 1
    start = next(i for i, r in enumerate(rows) if _width(r) >= needed)
    skipped = [' '.join(str(_cell(c)) for c in r if _cell(c) is not None) for r in rows[:start]]
    header_row, data = rows[start], rows[start + 1:]
    width = max(len(r) for r in [header_row, *data])
    keep = [i for i in range(width) if (i < len(header_row) and _cell(header_row[i]) is not None)
            or any(i < len(r) and _cell(r[i]) is not None for r in data)]
    header = _headers([header_row[i] if i < len(header_row) else None for i in keep])
    records = [{name: _cell(row[i]) if i < len(row) else None for i, name in zip(keep, header)} for row in data]
    return records, skipped


def _decode(data: bytes) -> str:
    for encoding in ('utf-8-sig', 'gb18030'):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise SourceFileError('CSV 文件不是 UTF-8 或 GBK 编码，请另存为 UTF-8 后再上传')


def _read_xlsx(data: bytes) -> dict[str, list[dict]]:
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise SourceFileError(f'无法作为 Excel（.xlsx）读取：{exc}') from None
    try:
        return {sheet.title: _records([list(r) for r in sheet.iter_rows(values_only=True)]) for sheet in book.worksheets}
    finally:
        book.close()


def _sheets(filename: str, data: bytes) -> dict:
    """One uploaded table file to its sheets; a CSV is one sheet named after the file."""
    suffix = PurePath(filename).suffix.lower()
    if suffix not in TABLE_SUFFIXES:
        raise SourceFileError(f'数据表只支持 {" / ".join(TABLE_SUFFIXES)}，不支持 {suffix or "无扩展名"} 文件')
    if not data:
        raise SourceFileError('文件是空的')
    if suffix == '.csv':
        return {PurePath(filename).stem: _records(list(csv.reader(io.StringIO(_decode(data)))))}
    return _read_xlsx(data)


def _bundle(sources: dict, file: dict, purpose: str | None) -> dict:
    if not sources:
        raise SourceFileError('文件里没有数据行：第一行应是表头，下面是数据')
    return {
        'schema': 'company_source_bundle.v1',
        'decision': (purpose or '').strip() or DEFAULT_PURPOSE,
        'file': file,
        'sources': sources,
    }


def load_table_file(filename: str, data: bytes, purpose: str | None = None) -> dict:
    if len(data) > MAX_BYTES:
        raise SourceFileError(f'文件太大（{len(data) // 1024 // 1024} MB），上限 {MAX_BYTES // 1024 // 1024} MB')
    sources = {name: {'records': records, 'requests': [], **({'skipped_rows': skipped} if skipped else {})}
               for name, (records, skipped) in _sheets(filename, data).items() if records}
    return _bundle(sources, {'name': PurePath(filename).name, 'kind': 'table', 'sha256': hashlib.sha256(data).hexdigest()}, purpose)


def load_table_files(files: list[tuple[str, bytes]], purpose: str | None = None) -> dict:
    """Several tables uploaded together, read as one set so the model can see how they relate. A table keeps its own
    name; only a real collision inside this batch gets the file's name added, because that name is what later runs
    match a confirmed ontology by."""
    if not files:
        raise SourceFileError('没有选择文件')
    if len(files) == 1:
        return load_table_file(files[0][0], files[0][1], purpose)
    total = sum(len(data) for _, data in files)
    if total > MAX_BYTES:
        raise SourceFileError(f'这些文件一共 {total // 1024 // 1024} MB，上限 {MAX_BYTES // 1024 // 1024} MB，请分批上传')
    sources: dict = {}
    for filename, data in files:
        try:
            sheets = _sheets(filename, data)
        except SourceFileError as exc:
            raise SourceFileError(f'{PurePath(filename).name}：{exc}') from None
        for name, (records, skipped) in sheets.items():
            if not records:
                continue
            key = name
            if key in sources:
                stem = PurePath(filename).stem
                named = f'{name}（{stem}）'
                key = named if stem != name and named not in sources else next(
                    f'{name}（{n}）' for n in range(2, 99) if f'{name}（{n}）' not in sources)
            sources[key] = {'records': records, 'requests': [], **({'skipped_rows': skipped} if skipped else {})}
    digests = sorted(hashlib.sha256(data).hexdigest() for _, data in files)
    file = {'name': '、'.join(PurePath(f).name for f, _ in files)[:200], 'kind': 'table',
            'sha256': hashlib.sha256(''.join(digests).encode()).hexdigest(),
            'files': [{'name': PurePath(f).name, 'sha256': hashlib.sha256(d).hexdigest()} for f, d in files]}
    return _bundle(sources, file, purpose)
