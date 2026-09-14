"""Write the synthetic demo company workbook (examples/company/demo_company.xlsx). Deterministic; all data is made up.

Two data problems are planted on purpose so the evaluation has something to find:
customer KH012 appears twice with different cities, and three orders name customer KH099, who is not in the customer table.
"""
import argparse
from datetime import date, datetime, timedelta
from pathlib import Path
import random

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
CITIES = ['苏州', '无锡', '宁波', '合肥', '东莞']
INDUSTRIES = ['电子', '汽车零部件', '家电', '医疗器械']
PRODUCTS = [('伺服电机 S-200', '电机', 1800), ('伺服电机 S-400', '电机', 2600), ('步进电机 B-57', '电机', 420),
            ('运动控制器 MC-8', '控制器', 5200), ('PLC 模块 P-16', '控制器', 3100), ('变频器 V-5.5', '控制器', 2300),
            ('光电传感器 G-10', '传感器', 160), ('接近开关 J-12', '传感器', 90), ('编码器 E-1024', '传感器', 680),
            ('温度传感器 T-100', '传感器', 210), ('减速机 R-60', '传动', 1500), ('直线模组 L-300', '传动', 3900)]
ISSUES = {'安装调试': '现场安装后{p}无法正常启动，需要工程师上门调试参数。',
          '质量缺陷': '{p}运行约两周后出现异常发热和报警，客户要求检测并更换。',
          '物流破损': '收货时{p}外包装破损，设备外壳有磕碰痕迹。',
          '使用咨询': '客户咨询{p}与现有产线控制系统的接线和通讯方式。'}
ENGINEERS = ['售后工程师-甲', '售后工程师-乙', '售后工程师-丙', '售后工程师-丁']


def build(rng: random.Random) -> dict[str, list[list]]:
    customers = [[f'KH{n:03d}', f'{rng.choice(CITIES)}{rng.choice(INDUSTRIES)}客户{n:02d}', rng.choice(INDUSTRIES),
                  rng.choice(CITIES), rng.choice('ABC')] for n in range(1, 31)]
    for row in customers:  # keep the name's city consistent with the city column
        row[1] = f'{row[3]}{row[2]}客户{row[0][-2:]}'
    dup = list(customers[11])
    dup[3] = next(c for c in CITIES if c != dup[3])
    customers.insert(12, dup)                                   # planted: KH012 twice, different cities
    products = [[f'CP{n:02d}', name, kind, price] for n, (name, kind, price) in enumerate(PRODUCTS, start=1)]
    orders = []
    for n in range(1, 161):
        customer = 'KH099' if n in (40, 90, 140) else f'KH{rng.randint(1, 30):03d}'   # planted: unknown customer
        product = rng.choice(products)
        qty = rng.randint(1, 20)
        orders.append([f'DD2026{n:04d}', customer, product[0], qty, qty * product[3],
                       date(2026, 1, 5) + timedelta(days=rng.randint(0, 220)), rng.choice(['已交付'] * 6 + ['生产中'] * 3 + ['已取消'])])
    tickets = []
    for n in range(1, 46):
        order = rng.choice([o for o in orders if o[6] == '已交付'])
        issue = rng.choice(list(ISSUES))
        product_name = next(p[1] for p in products if p[0] == order[2])
        tickets.append([f'SH{n:04d}', order[0], issue, ISSUES[issue].format(p=product_name),
                        order[5] + timedelta(days=rng.randint(3, 60)), rng.choice(['已关闭'] * 3 + ['处理中']), rng.choice(ENGINEERS)])
    return {
        '客户': [['客户编号', '客户名称', '行业', '城市', '客户等级'], *customers],
        '产品': [['产品编号', '产品名称', '品类', '单价'], *products],
        '订单': [['订单号', '客户编号', '产品编号', '数量', '金额', '下单日期', '订单状态'], *orders],
        '售后工单': [['工单号', '订单号', '问题类型', '问题描述', '创建日期', '处理状态', '负责工程师'], *tickets],
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT / 'examples/company/demo_company.xlsx')
    args = parser.parse_args(argv)
    book = Workbook()
    book.remove(book.active)
    for name, rows in build(random.Random(20260914)).items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book.properties.creator = 'OntoPoc demo (synthetic data)'
    book.properties.created = book.properties.modified = datetime(2026, 9, 14)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    book.save(args.output)
    print(f'wrote {args.output.name}: ' + ', '.join(f'{s.title} {s.max_row - 1} rows' for s in book.worksheets))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
